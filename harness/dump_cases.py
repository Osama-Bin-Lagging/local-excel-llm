"""Export Formula Tests + their context ranges from the test workbook to JSON, as Excel would pass them."""
import datetime as dt, json, re, sys
from openpyxl import load_workbook
from openpyxl.utils import range_boundaries

SRC = sys.argv[1]
wb = load_workbook(SRC)
EPOCH = dt.datetime(1899, 12, 30)


def cell_value(ws, r, c):
    v = ws.cell(row=r, column=c).value
    if isinstance(v, dt.datetime):
        return (v - EPOCH).days  # Excel passes dates as serial numbers
    if isinstance(v, str) and v.startswith("=") and ws.title == "Sales" and c == 6:
        return ws.cell(row=r, column=4).value * ws.cell(row=r, column=5).value  # Revenue = Units * Price
    return "" if v is None else v


def read_range(ref):
    sheet, rng = ref.split("!")
    ws = wb[sheet]
    c1, r1, c2, r2 = range_boundaries(rng)
    return [[cell_value(ws, r, c) for c in range(c1, c2 + 1)] for r in range(r1, r2 + 1)]


ft = wb["Formula Tests"]
cases = []
for r in range(2, ft.max_row + 1):
    fid = ft.cell(row=r, column=1).value
    if not (fid and re.match(r"F\d\d", str(fid))):
        continue
    ctx = ft.cell(row=r, column=4).value
    fn = ft.cell(row=r, column=5).value
    cases.append({
        "id": fid, "row": r, "skill": ft.cell(row=r, column=2).value,
        "prompt": ft.cell(row=r, column=3).value,
        "fn": "ask" if fn.startswith(("=LOKUP.H", "=LLAMA.ASK")) else "formula",
        "contextAddress": None if ctx == "(none)" else ctx,
        "context": None if ctx == "(none)" else read_range(ctx),
        "parts": ft.cell(row=r, column=8).value, "expect": ft.cell(row=r, column=9).value,
    })
json.dump(cases, sys.stdout, indent=1)
