"""Held-out v4: a fresh business workbook (not derived from any test file) to check that the
glossary generalises. Builds the answer-key workbook, has Excel compute the references, writes tasks.

Usage: python3 -I build_v4.py <out tasks_v4.json> <keys dir>
"""
import json, os, shutil, subprocess, sys, tempfile
from openpyxl import Workbook
from openpyxl.utils import range_boundaries, get_column_letter

OUT, KEYDIR = sys.argv[1], sys.argv[2]
SEP = "\x1f"
wb = Workbook()

# ---- SaaS Metrics (title row 1, header row 3, data rows 4-9)
ws = wb.active
ws.title = "SaaS Metrics"
ws["A1"] = "SaaS monthly metrics (INR)"
ws.append([])
ws.append(["Month", "Customers at start", "New customers", "Customers lost", "MRR (INR)", "Sales & marketing spend",
           "Website visitors", "Signups", "Churn rate", "CAC", "Conversion rate", "MRR growth", "ARPU", "ARR"])
data = [("Jan", 1200, 150, 36, 1800000, 900000, 52000, 1560), ("Feb", 1314, 170, 42, 1990000, 980000, 55000, 1650),
        ("Mar", 1442, 160, 50, 2150000, 1050000, 58500, 1700), ("Apr", 1552, 190, 47, 2330000, 1120000, 61000, 1830),
        ("May", 1695, 210, 55, 2560000, 1200000, 64000, 1980), ("Jun", 1850, 200, 61, 2760000, 1180000, 66500, 2050)]
for i, row in enumerate(data, 4):
    ws.append(list(row) + [f"=D{i}/B{i}", f"=F{i}/C{i}", f"=H{i}/G{i}", f"=E{i}/E{i-1}-1" if i > 4 else None, f"=E{i}/B{i}", f"=E{i}*12"])

# ---- PnL (years across, header row 1)
ws = wb.create_sheet("PnL")
ws.append(["INR mm", "FY2022", "FY2023", "FY2024", "FY2025"])
for label, vals in [("Revenue", [8000, 9200, 10500, 11800]), ("COGS", [4800, 5400, 6050, 6700]),
                    ("Operating expenses (excl. D&A)", [1500, 1650, 1800, 1950]), ("D&A", [300, 330, 360, 390]),
                    ("Interest expense", [200, 190, 180, 170]), ("Tax rate", [0.25] * 4)]:
    ws.append([label] + vals)
ws["A8"], ws["A9"], ws["A10"], ws["A11"], ws["A12"] = "Gross margin", "EBIT", "Net income", "Interest coverage", "Revenue CAGR FY2022-FY2025"
for c in "BCDE":
    ws[f"{c}8"] = f"=({c}2-{c}3)/{c}2"
    ws[f"{c}9"] = f"={c}2-{c}3-{c}4-{c}5"
    ws[f"{c}10"] = f"=({c}9-{c}6)*(1-{c}7)"
    ws[f"{c}11"] = f"={c}9/{c}6"
ws["B12"] = "=(E2/B2)^(1/3)-1"

# ---- Loan Calc
ws = wb.create_sheet("Loan Calc")
for r, (a, b) in enumerate([("Loan amount (INR)", 2500000), ("Annual interest rate", 0.09), ("Tenure (years)", 20)], 2):
    ws[f"A{r}"], ws[f"B{r}"] = a, b
ws["A1"] = "Home loan"
ws["A6"], ws["B6"] = "Monthly EMI", "=PMT(B3/12,B4*12,-B2)"
ws["A7"], ws["B7"] = "Total interest paid", "=B6*B4*12-B2"

# ---- Budget
ws = wb.create_sheet("Budget")
ws.append(["Department", "Budget", "Actual", "Variance", "Variance %"])
for i, (d, b, a) in enumerate([("Marketing", 500000, 565000), ("Sales", 800000, 760000), ("IT", 650000, 702000),
                               ("HR", 300000, 288000), ("Operations", 900000, 945000), ("Finance", 250000, 240000)], 2):
    ws.append([d, b, a, f"=C{i}-B{i}", f"=C{i}/B{i}-1"])

# ---- Stock Data
ws = wb.create_sheet("Stock Data")
ws.append(["Company", "Share price", "Shares (mm)", "Net income", "Debt", "Cash", "Market cap", "EV", "P/E"])
for i, (n, p, s, ni, d, c) in enumerate([("Alpha Ltd", 120, 50, 300, 1500, 400), ("Beta Ltd", 450, 20, 720, 800, 350),
                                         ("Gamma Ltd", 88, 120, 580, 2400, 600), ("Delta Ltd", 1500, 8, 1100, 500, 900),
                                         ("Epsilon Ltd", 310, 35, 640, 900, 200)], 2):
    ws.append([n, p, s, ni, d, c, f"=B{i}*C{i}", f"=G{i}+E{i}-F{i}", f"=G{i}/D{i}"])

# ---- Inventory
ws = wb.create_sheet("Inventory")
ws.append(["Item", "COGS", "Opening inventory", "Closing inventory", "Inventory turnover"])
for i, (n, cg, o, c) in enumerate([("Widgets", 4800000, 600000, 680000), ("Gadgets", 3100000, 520000, 460000),
                                   ("Gizmos", 2650000, 300000, 362000), ("Doohickeys", 1900000, 410000, 390000)], 2):
    ws.append([n, cg, o, c, f"=B{i}/((C{i}+D{i})/2)"])

# ---- Project
ws = wb.create_sheet("Project")
ws.append(["Year", 0, 1, 2, 3, 4, 5])
ws.append(["Cash flow (INR)", -10000000, 2500000, 3000000, 3200000, 3500000, 2800000])
ws["A4"], ws["B4"] = "Discount rate", 0.12
ws["A6"], ws["B6"] = "NPV", "=NPV(B4,C2:G2)+B2"
ws["A7"], ws["B7"] = "IRR", "=IRR(B2:G2)"

# ---- Scores
ws = wb.create_sheet("Scores")
ws.append(["Student", "Score", "z-score", "Rank"])
for i, (n, sc) in enumerate(zip(["Asha", "Bilal", "Chen", "Divya", "Elan", "Farah", "Gopal", "Hina", "Ishan", "Jaya"],
                                [72, 85, 64, 90, 78, 55, 88, 69, 94, 81]), 2):
    ws.append([n, sc, f"=(B{i}-AVERAGE($B$2:$B$11))/STDEV($B$2:$B$11)", f"=RANK(B{i},$B$2:$B$11,0)"])

key = os.path.join(KEYDIR, "V4_key.xlsx")
wb.save(key)

S, P, L, B, K, I, J, C = ("'SaaS Metrics'!A3:N9", "'PnL'!A1:E12", "'Loan Calc'!A1:B7", "'Budget'!A1:E7",
                          "'Stock Data'!A1:I6", "'Inventory'!A1:E5", "'Project'!A1:G7", "'Scores'!A1:D11")
tasks = [
    ("SaaS Metrics", "I7", "Churn rate for April (row 7).", S),
    ("SaaS Metrics", "J6", "CAC for March (row 6).", S),
    ("SaaS Metrics", "K9", "Visitor-to-signup conversion rate for June (row 9).", S),
    ("SaaS Metrics", "L8", "MRR growth for May versus April (row 8).", S),
    ("SaaS Metrics", "M5", "ARPU for February (row 5), per customer at the start of the month.", S),
    ("SaaS Metrics", "N9", "ARR for June (row 9).", S),
    ("PnL", "D8", "Gross margin for FY2024.", P),
    ("PnL", "E9", "EBIT for FY2025.", P),
    ("PnL", "C10", "Net income for FY2023.", P),
    ("PnL", "E11", "Interest coverage for FY2025.", P),
    ("PnL", "B12", "Revenue CAGR from FY2022 to FY2025.", P),
    ("Loan Calc", "B6", "Monthly EMI for this loan.", L),
    ("Loan Calc", "B7", "Total interest paid over the life of the loan (the EMI is in B6).", L),
    ("Budget", "D4", "Variance (actual minus budget) for IT, row 4.", B),
    ("Budget", "E6", "Variance % versus budget for Operations, row 6.", B),
    ("Stock Data", "G3", "Market cap for Beta Ltd (row 3).", K),
    ("Stock Data", "H5", "Enterprise value for Delta Ltd (row 5).", K),
    ("Stock Data", "I2", "P/E ratio for Alpha Ltd (row 2).", K),
    ("Inventory", "E4", "Inventory turnover for Gizmos (row 4).", I),
    ("Project", "B6", "NPV of the project at the discount rate in B4.", J),
    ("Project", "B7", "IRR of the project.", J),
    ("Scores", "C5", "z-score of Divya's score (row 5) against the whole class.", C),
    ("Scores", "D7", "Rank of Farah's score (row 7) in the class, highest score = 1.", C),
]


def q(s):
    return '"' + s.replace("\\", "\\\\").replace('"', '\\"') + '"'


# Excel computes every cell we need (targets + formula cells inside each context)
from openpyxl import load_workbook
kf = load_workbook(key)
wanted = {}
for sh, cell, _, ctx in tasks:
    csh, rng = ctx.split("!")
    csh = csh.strip("'")
    c1, r1, c2, r2 = range_boundaries(rng)
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            wanted.setdefault(csh, set()).add(f"{get_column_letter(c)}{r}")
name = f"ref_{os.getpid()}_V4.xlsx"
tmp = os.path.join(tempfile.mkdtemp(), name)
shutil.copy(key, tmp)
L_ = ["with timeout of 180 seconds", 'tell application "Microsoft Excel"', f"open POSIX file {q(tmp)}",
      f"repeat 20 times\ntry\nset wb to workbook {q(name)}\nexit repeat\non error\ndelay 0.5\nend try\nend repeat",
      "calculate", 'set acc to ""']
for sh, cells in wanted.items():
    for a in sorted(cells):
        L_.append(f'set acc to acc & {q(sh + "!" + a)} & {q(SEP)} & ((value of range "{a}" of worksheet {q(sh)} of wb) as text) & linefeed')
L_ += ["close wb saving no", "return acc", "end tell", "end timeout"]
out = subprocess.run(["osascript", "-"], input="\n".join(L_), capture_output=True, text=True)
if out.returncode:
    sys.exit(out.stderr)
vals = {}
for line in out.stdout.splitlines():
    if SEP in line:
        k_, v = line.split(SEP, 1)
        vals[tuple(k_.split("!", 1))] = v


def val(v):
    if v in ("", "missing value"):
        return ""
    try:
        return float(v.replace(",", ""))
    except ValueError:
        return v


result = []
for sh, cell, prompt, ctx in tasks:
    csh, rng = ctx.split("!")
    csh = csh.strip("'")
    c1, r1, c2, r2 = range_boundaries(rng)
    grid = [["" if (csh == sh and f"{get_column_letter(c)}{r}" == cell) else val(vals[(csh, f"{get_column_letter(c)}{r}")])
             for c in range(c1, c2 + 1)] for r in range(r1, r2 + 1)]
    result.append({"test": "V4", "sheet": sh, "cell": cell, "prompt": prompt, "contextAddress": ctx, "context": grid,
                   "ref": val(vals[(sh, cell)]), "key_formula": kf[sh][cell].value, "address": f"'{sh}'!{cell}",
                   "id": f"V4-{sh.split()[0][:5]}-{cell}", "fn": "formula"})
json.dump(result, open(OUT, "w"), indent=1)
print(len(result), "tasks")
for t in result:
    print(f"  {t['id']:<16} ref={str(t['ref'])[:14]:<14} key={t['key_formula']}")
