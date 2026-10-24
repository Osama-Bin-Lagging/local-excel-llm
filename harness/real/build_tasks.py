"""Turn the two downloaded Excel assessments into LOKUP.V tasks with Excel-computed reference values.

Each task = one cell a candidate must fill. The prompt uses only candidate-visible text (row/column
labels, the sheet's guidance column, Read Me task text) — never the answer key's hints.
Reference values come from the answer key, calculated by real Excel.

Usage: python3 -I build_tasks.py <dir with the 4 xlsx> <out tasks.json> <out patched-key dir>
"""
import datetime as dt, json, os, shutil, subprocess, sys, tempfile
from openpyxl import load_workbook
from openpyxl.utils import range_boundaries, get_column_letter

SRC, OUT, KEYDIR = sys.argv[1], sys.argv[2], sys.argv[3]
EPOCH = dt.datetime(1899, 12, 30)
SEP = "\x1f"
os.makedirs(KEYDIR, exist_ok=True)

# ---- answer-key fix: Trading Comps!B17 is empty in the key but Summary!B5 links to it.
# The key's own notes say "MEDIAN of peer EV / EBITDA multiples", so we fill it with exactly that.
ib = load_workbook(os.path.join(SRC, "IB_Excel_Test_Answer_Key.xlsx"))
ib["Trading Comps"]["B17"] = "=MEDIAN(K5:K14)"
ib.save(os.path.join(KEYDIR, "IB_key.xlsx"))
shutil.copy(os.path.join(SRC, "Data_Analytics_Excel_Test_Answer_Key.xlsx"), os.path.join(KEYDIR, "DA_key.xlsx"))

OD, TC, DM = "'Operating Data'!A4:E13", "'Trading Comps'!A4:L24", "'Deal Model'!A4:E29"
RD, AN = "'Raw Data'!A4:R124", "'Analysis'!A4:C17"

tasks = []
def T(test, sheet, cell, prompt, ctx):
    tasks.append({"test": test, "sheet": sheet, "cell": cell, "prompt": prompt, "contextAddress": ctx})

# ---------------- IB
T("IB", "Operating Data", "C12", "Revenue growth for AcquirerCo FY2024E vs FY2023A.", OD)
T("IB", "Operating Data", "E12", "Revenue growth for TargetCo FY2024E vs FY2023A.", OD)
T("IB", "Operating Data", "B13", "EBITDA margin for AcquirerCo FY2023A.", OD)
T("IB", "Operating Data", "E13", "EBITDA margin for TargetCo FY2024E.", OD)
for cell, what, co in [("H5", "Equity Value", "Aster Consumer Ltd (row 5)"), ("I5", "Enterprise Value", "Aster Consumer Ltd (row 5)"),
                       ("J5", "EV / FY2023A Revenue", "Aster Consumer Ltd (row 5)"), ("K5", "EV / FY2023A EBITDA", "Aster Consumer Ltd (row 5)"),
                       ("H10", "Equity Value", "Fable Retail Ltd (row 10)"), ("I10", "Enterprise Value", "Fable Retail Ltd (row 10)"),
                       ("K12", "EV / FY2023A EBITDA", "Harbor Foods Ltd (row 12)")]:
    T("IB", "Trading Comps", cell, f"{what} for {co}.", TC)
T("IB", "Trading Comps", "B17", "Peer median EV / EBITDA across the peer set (rows 5 to 14).", TC)
T("IB", "Trading Comps", "B21", "Implied Enterprise Value for TargetCo: apply the peer median EV / EBITDA multiple to TargetCo FY2023A EBITDA.", TC)
T("IB", "Trading Comps", "B24", "Implied Equity Value for TargetCo: implied enterprise value minus TargetCo net debt (debt less cash).", TC)
dm_cand = load_workbook(os.path.join(SRC, "IB_Excel_Test_Candidate.xlsx"))["Deal Model"]
for r in range(18, 30):
    T("IB", "Deal Model", f"B{r}", f"{dm_cand[f'A{r}'].value}: {dm_cand[f'D{r}'].value}.", DM)
for r, src in [(5, TC), (6, TC), (7, TC), (8, DM), (9, DM), (10, DM), (11, DM), (12, DM)]:
    label = load_workbook(os.path.join(SRC, "IB_Excel_Test_Candidate.xlsx"))["Summary"][f"A{r}"].value
    T("IB", "Summary", f"B{r}", f"Link the {label} from the model (just a link to the cell that holds it).", src)

# ---------------- Data Analytics
da = load_workbook(os.path.join(SRC, "Data_Analytics_Excel_Test_Answer_Key.xlsx"))
raw = da["Raw Data"]
not_delivered = next(r for r in range(5, 125) if raw[f"M{r}"].value != "Delivered")
blank_disc = next(r for r in range(5, 125) if raw[f"J{r}"].value in (None, ""))
T("DA", "Raw Data", "O5", "Gross Sales (INR) for row 5: Unit Price × Quantity.", RD)
T("DA", "Raw Data", "P5", "Net Sales (INR) for row 5: Gross Sales after the discount; treat a missing discount as 0%.", RD)
T("DA", "Raw Data", f"P{blank_disc}", f"Net Sales (INR) for row {blank_disc}: Gross Sales after the discount; treat a missing discount as 0%.", RD)
T("DA", "Raw Data", "Q5", "Gross Profit (INR) for row 5: Net Sales − Cost.", RD)
T("DA", "Raw Data", "R5", "Delivered Flag for row 5: 1 if Status is Delivered, otherwise 0.", RD)
T("DA", "Raw Data", f"R{not_delivered}", f"Delivered Flag for row {not_delivered}: 1 if Status is Delivered, otherwise 0.", RD)
an_cand = load_workbook(os.path.join(SRC, "Data_Analytics_Excel_Test_Candidate.xlsx"))["Analysis"]
for r in range(5, 18):
    ctx = AN if r in (12, 13) else RD  # gross margin / AOV are built from other KPIs on Analysis
    T("DA", "Analysis", f"B{r}", f"{an_cand[f'A{r}'].value}: {an_cand[f'C{r}'].value}.", ctx)
for cell, seg_cell, seg, col in [("F6", "E6", "North region", "Net Sales"), ("G6", "E6", "North region", "Gross Profit"),
                                 ("F13", "E13", "Website channel", "Net Sales"), ("F18", "E18", "Electronics category", "Net Sales")]:
    T("DA", "Analysis", cell, f"{col} (INR) for the {seg} (segment name is in {seg_cell}), delivered orders only.", RD)

# ---------------- reference + context values via real Excel
def q(s):
    return '"' + s.replace("\\", "\\\\").replace('"', '\\"') + '"'

def excel_values(path, wanted):
    """wanted: {sheet: [range...]} → {(sheet, coord): text}. Excel computes every formula."""
    # sandboxed Excel opens files from the temp folder without a permission prompt
    name = f"ref_{os.getpid()}_{os.path.basename(path)}"
    tmp_path = os.path.join(tempfile.mkdtemp(), name)
    shutil.copy(path, tmp_path)
    path = tmp_path
    lines = ["with timeout of 180 seconds", 'tell application "Microsoft Excel"', f"open POSIX file {q(path)}",
             f"repeat 20 times\ntry\nset wb to workbook {q(name)}\nexit repeat\non error\ndelay 0.5\nend try\nend repeat",
             "calculate", 'set acc to ""']
    for sheet, ranges in wanted.items():
        for rng in ranges:
            c1, r1, c2, r2 = range_boundaries(rng)
            for r in range(r1, r2 + 1):
                for c in range(c1, c2 + 1):
                    a = f"{get_column_letter(c)}{r}"
                    lines.append(f'set acc to acc & {q(sheet + "!" + a)} & {q(SEP)} & (string value of range "{a}" of worksheet {q(sheet)} of wb) & {q(SEP)} & ((value of range "{a}" of worksheet {q(sheet)} of wb) as text) & linefeed')
    lines += ["close wb saving no", "return acc", "end tell", "end timeout"]
    out = subprocess.run(["osascript", "-"], input="\n".join(lines), capture_output=True, text=True)
    if out.returncode:
        sys.exit(out.stderr)
    vals = {}
    for line in out.stdout.splitlines():
        if not line.strip():
            continue
        key, shown, raw_v = (line.split(SEP) + ["", ""])[:3]
        sh, a = key.split("!", 1)
        vals[(sh, a)] = (shown, raw_v)
    return vals

def as_value(raw_v, shown):
    if shown.startswith("#"):
        return shown
    if raw_v in ("", "missing value"):
        return ""
    try:
        return float(raw_v.replace(",", "")) if any(ch.isdigit() for ch in raw_v) and not any(ch.isalpha() and ch not in "Ee" for ch in raw_v) else raw_v
    except ValueError:
        return raw_v

for test, path in [("IB", os.path.join(KEYDIR, "IB_key.xlsx")), ("DA", os.path.join(KEYDIR, "DA_key.xlsx"))]:
    wb_f = load_workbook(path)
    mine = [t for t in tasks if t["test"] == test]
    # formula cells inside any context range + every target → ask Excel for those only
    wanted = {}
    for t in mine:
        sh, rng = t["contextAddress"].split("!")
        sh = sh.strip("'")
        c1, r1, c2, r2 = range_boundaries(rng)
        cells = [f"{get_column_letter(c)}{r}" for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)
                 if wb_f[sh][f"{get_column_letter(c)}{r}"].data_type == "f"]
        wanted.setdefault(sh, set()).update(cells)
        wanted.setdefault(t["sheet"], set()).add(t["cell"])
    vals = excel_values(path, {sh: sorted(cs) for sh, cs in wanted.items()})
    for t in mine:
        t["ref"] = as_value(*vals[(t["sheet"], t["cell"])][::-1])
        t["ref_shown"] = vals[(t["sheet"], t["cell"])][0]
        t["key_formula"] = wb_f[t["sheet"]][t["cell"]].value
        sh, rng = t["contextAddress"].split("!")
        sh = sh.strip("'")
        c1, r1, c2, r2 = range_boundaries(rng)
        grid = []
        for r in range(r1, r2 + 1):
            row = []
            for c in range(c1, c2 + 1):
                a = f"{get_column_letter(c)}{r}"
                cell = wb_f[sh][a]
                if sh == t["sheet"] and a == t["cell"]:
                    row.append("")  # the cell being asked about is still empty for the user
                elif cell.data_type == "f":
                    row.append(as_value(*vals[(sh, a)][::-1]))
                elif isinstance(cell.value, dt.datetime):
                    row.append((cell.value - EPOCH).days)
                else:
                    row.append("" if cell.value is None else cell.value)
            grid.append(row)
        t["context"] = grid
        t["address"] = f"'{t['sheet']}'!{t['cell']}"
        t["id"] = f"{test}-{t['sheet'].split()[0][:4]}-{t['cell']}"
        t["fn"] = "formula"

json.dump(tasks, open(OUT, "w"), indent=1, default=str)
print(len(tasks), "tasks →", OUT)
for t in tasks:
    print(f"  {t['id']:<22} ref={str(t['ref_shown'])[:14]:<14} key={t['key_formula'][:50]}")
